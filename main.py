import io
import csv
import os
import time
import re
import tempfile
import subprocess
from datetime import date, datetime
from decimal import Decimal
from urllib.parse import urlparse

import psycopg2
from psycopg2 import errors
from psycopg2.extras import RealDictCursor
from flask import (
    Flask, jsonify, request, abort,
    send_file, send_from_directory,
)

from openpyxl import load_workbook, Workbook


# ============================================================
# CONFIG
# ============================================================
FRONTEND_DIST = os.environ.get("FRONTEND_DIST", "/app/frontend_dist")


# ============================================================
# SEED (встроен)
# ============================================================
def _wait_db(url: str, attempts: int = 30) -> None:
    for i in range(attempts):
        try:
            psycopg2.connect(url).close()
            return
        except psycopg2.OperationalError:
            print(f"[seed] waiting for DB ({i + 1}/{attempts})…")
            time.sleep(1)
    raise RuntimeError("DB not available")


def _as_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
            try:
                return datetime.strptime(v, fmt).date()
            except ValueError:
                pass
    return None


def _as_float(v) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _run_seed():
    url = os.environ["DATABASE_URL"]
    path = os.environ.get("EXCEL_PATH", "data/траты.xlsx")
    _wait_db(url)

    conn = psycopg2.connect(url)
    conn.autocommit = False
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT COUNT(*) FROM expenses")
            if cur.fetchone()[0] > 0:
                print("[seed] DB already populated — skip")
                return
            cur.execute("SELECT COUNT(*) FROM incomes")
            if cur.fetchone()[0] > 0:
                print("[seed] DB already populated — skip")
                return

        if not os.path.exists(path):
            print(f"[seed] Excel not found at {path} — start with empty DB")
            return

        print(f"[seed] reading {path}…")
        wb = load_workbook(path, data_only=True)

        ws = wb["Справочники"]
        cat_map, sub_map = {}, {}

        with conn.cursor() as cur:
            for c in range(1, 21):
                v = ws.cell(1, c).value
                if not v:
                    continue
                name = str(v).strip()
                if not name:
                    continue
                cur.execute(
                    "INSERT INTO expense_categories(name) VALUES(%s) "
                    "ON CONFLICT(name) DO UPDATE SET name=EXCLUDED.name RETURNING id",
                    (name,),
                )
                cid = cur.fetchone()[0]
                cat_map[name] = cid
                for r in range(2, ws.max_row + 1):
                    sv = ws.cell(r, c).value
                    if sv:
                        sub = str(sv).strip()
                        if sub:
                            cur.execute(
                                "INSERT INTO expense_subcategories(category_id, name) "
                                "VALUES(%s, %s) ON CONFLICT(category_id, name) DO NOTHING RETURNING id",
                                (cid, sub),
                            )
                            got = cur.fetchone()
                            if got:
                                sub_map[(cid, sub)] = got[0]
                            else:
                                cur.execute(
                                    "SELECT id FROM expense_subcategories WHERE category_id=%s AND name=%s",
                                    (cid, sub),
                                )
                                sub_map[(cid, sub)] = cur.fetchone()[0]

            for r in range(2, ws.max_row + 1):
                v = ws.cell(r, 21).value
                if v:
                    name = str(v).strip()
                    if name:
                        cur.execute(
                            "INSERT INTO income_categories(name) VALUES(%s) "
                            "ON CONFLICT(name) DO NOTHING",
                            (name,),
                        )
            conn.commit()
            cur.execute("SELECT id, name FROM income_categories")
            inc_cat_map = {name: cid for cid, name in cur.fetchall()}

        # Расходы
        ws = wb["Учет расходов"]
        with conn.cursor() as cur:
            n = 0
            for r in range(2, ws.max_row + 1):
                d = _as_date(ws.cell(r, 3).value)
                raw_cat = ws.cell(r, 4).value
                if not d or not raw_cat:
                    continue
                cat = str(raw_cat).strip()
                if not cat:
                    continue
                cid = cat_map.get(cat)
                if not cid:
                    cur.execute(
                        "INSERT INTO expense_categories(name) VALUES(%s) RETURNING id",
                        (cat,),
                    )
                    cid = cur.fetchone()[0]
                    cat_map[cat] = cid

                sub = ws.cell(r, 5).value
                sub = str(sub).strip() if sub else None
                sid = None
                if sub:
                    key = (cid, sub)
                    if key in sub_map:
                        sid = sub_map[key]
                    else:
                        cur.execute(
                            "INSERT INTO expense_subcategories(category_id, name) "
                            "VALUES(%s, %s) RETURNING id",
                            (cid, sub),
                        )
                        sid = cur.fetchone()[0]
                        sub_map[key] = sid

                cur.execute(
                    "INSERT INTO expenses(date, category_id, subcategory_id, description, amount, comment) "
                    "VALUES(%s, %s, %s, %s, %s, %s)",
                    (d, cid, sid,
                     (str(ws.cell(r, 6).value).strip() if ws.cell(r, 6).value else None),
                     _as_float(ws.cell(r, 7).value),
                     (str(ws.cell(r, 8).value).strip() if ws.cell(r, 8).value else None)),
                )
                n += 1
            conn.commit()
            print(f"[seed] imported {n} expenses")

        # Доходы
        ws = wb["Учет доходов"]
        with conn.cursor() as cur:
            n = 0
            for r in range(2, ws.max_row + 1):
                d = _as_date(ws.cell(r, 3).value)
                raw_cat = ws.cell(r, 4).value
                if not d or not raw_cat:
                    continue
                cat = str(raw_cat).strip()
                if not cat:
                    continue
                cid = inc_cat_map.get(cat)
                if not cid:
                    cur.execute(
                        "INSERT INTO income_categories(name) VALUES(%s) RETURNING id",
                        (cat,),
                    )
                    cid = cur.fetchone()[0]
                    inc_cat_map[cat] = cid
                cur.execute(
                    "INSERT INTO incomes(date, category_id, description, amount, comment) "
                    "VALUES(%s, %s, %s, %s, %s)",
                    (d, cid,
                     (str(ws.cell(r, 5).value).strip() if ws.cell(r, 5).value else None),
                     _as_float(ws.cell(r, 6).value),
                     (str(ws.cell(r, 9).value).strip() if ws.cell(r, 9).value else None)),
                )
                n += 1
            conn.commit()
            print(f"[seed] imported {n} incomes")
        print("[seed] done ✔")
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


_run_seed()


# ============================================================
# APP (создаём ДО объявления любых роутов!)
# ============================================================
app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 32 * 1024 * 1024


# ---------------- CORS ----------------
@app.after_request
def _cors(resp):
    resp.headers["Access-Control-Allow-Origin"] = "*"
    resp.headers["Access-Control-Allow-Methods"] = "GET, POST, PUT, DELETE, OPTIONS"
    resp.headers["Access-Control-Allow-Headers"] = "Content-Type"
    return resp


@app.route("/api/<path:_any>", methods=["OPTIONS"])
def _cors_preflight(_any):
    return "", 204


# -------------------- helpers --------------------
def get_conn():
    return psycopg2.connect(os.environ["DATABASE_URL"], cursor_factory=RealDictCursor)


def q(sql, params=None, fetch="all"):
    conn = get_conn()
    try:
        with conn.cursor() as cur:
            cur.execute(sql, params or ())
            if cur.description:
                rows = cur.fetchall() if fetch == "all" else cur.fetchone()
            else:
                rows = None
        conn.commit()
        return rows
    finally:
        conn.close()


def ser(rows):
    def conv(r):
        d = dict(r)
        for k, v in list(d.items()):
            if isinstance(v, Decimal):
                d[k] = float(v)
            elif isinstance(v, date):
                d[k] = v.isoformat()
        return d

    if isinstance(rows, list):
        return [conv(r) for r in rows]
    return conv(rows) if rows else rows


def _parse_months_param():
    raw = request.args.get("months", "").strip()
    if not raw:
        return []
    try:
        return [int(x) for x in raw.split(",") if x.strip()]
    except ValueError:
        abort(400, "Неверный параметр months")

# ============================================================
# HELPERS для фильтрации/сортировки/пагинации
# ============================================================
def _parse_int(v, default=None):
    if v in (None, "", "null"):
        return default
    try:
        return int(v)
    except (TypeError, ValueError):
        return default


def _parse_float(v, default=None):
    if v in (None, "", "null"):
        return default
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return default


def _like(v: str) -> str:
    """Экранирует % и _ для ILIKE."""
    return "%" + v.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_") + "%"

def _escape_like(v: str) -> str:
    """Экранирует %, _ и \\ для использования в ILIKE ... ESCAPE '\\'."""
    return (
        v.replace("\\", "\\\\")
         .replace("%", "\\%")
         .replace("_", "\\_")
    )

def _apply_date_filter(where: list, p: list, expr: str,
                       date_from: str, date_to: str, date_eq: str = ""):
    """Умная фильтрация по дате.

    Приоритет:
      1) date_eq (одиночный день)      → date = ...
      2) from и to                     → BETWEEN / диапазон
      3) только from (без to)          → ТОЧНОЕ совпадение date = from
                                          (или >= from, если from неполное: месяц/год)
      4) только to (без from)          → date <= to
      5) оба пусты                     → без фильтра
    """
    def expand(v: str, is_end: bool):
        v = v.strip()
        if not v:
            return None
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
            return ("eq", v)
        if re.fullmatch(r"\d{4}-\d{2}", v):
            y, m = v.split("-")
            if is_end:
                nm_y, nm_m = (int(y), int(m) + 1) if int(m) < 12 else (int(y) + 1, 1)
                return ("lt", f"{nm_y:04d}-{nm_m:02d}-01")
            return ("gte", f"{v}-01")
        if re.fullmatch(r"\d{4}", v):
            if is_end:
                return ("lt", f"{int(v) + 1:04d}-01-01")
            return ("gte", f"{v}-01-01")
        return ("eq", v[:10])

    if date_eq:
        r = expand(date_eq, is_end=False)
        if r and r[0] == "eq":
            where.append(f"{expr} = %s"); p.append(r[1])
        elif r:
            op = ">=" if r[0] == "gte" else "="
            where.append(f"{expr} {op} %s"); p.append(r[1])
        return

    start = expand(date_from, is_end=False) if date_from else None
    end = expand(date_to, is_end=True) if date_to else None

    if start and end:
        # полноценный диапазон
        if start[0] == "eq":
            where.append(f"{expr} >= %s"); p.append(start[1])
        else:
            where.append(f"{expr} >= %s"); p.append(start[1])
        if end[0] == "lt":
            where.append(f"{expr} < %s"); p.append(end[1])
        else:
            where.append(f"{expr} <= %s"); p.append(end[1])
        return

    if start and not end:
        # только «От» — ТОЧНОЕ совпадение, если дата полная
        if start[0] == "eq":
            where.append(f"{expr} = %s"); p.append(start[1])
        else:
            where.append(f"{expr} >= %s"); p.append(start[1])
        return

    if end and not start:
        # только «До»
        if end[0] == "lt":
            where.append(f"{expr} < %s"); p.append(end[1])
        else:
            where.append(f"{expr} <= %s"); p.append(end[1])
        return


def _apply_amount_filter(where: list, p: list, expr: str,
                         amount_min, amount_max, amount_eq):
    """Точная фильтрация по сумме.

    Если задано точное значение — `amount = X`.
    Иначе — диапазон `>= / <=`.
    """
    if amount_eq is not None:
        where.append(f"{expr} = %s")
        p.append(amount_eq)
        return
    if amount_min is not None:
        where.append(f"{expr} >= %s")
        p.append(amount_min)
    if amount_max is not None:
        where.append(f"{expr} <= %s")
        p.append(amount_max)

# Разрешённые поля сортировки для каждого ресурса.
EXPENSE_SORTS = {
    "date": "e.date",
    "category": "ec.name",
    "subcategory": "es.name",
    "description": "e.description",
    "amount": "e.amount",
    "comment": "e.comment",
}
INCOME_SORTS = {
    "date": "i.date",
    "category": "ic.name",
    "description": "i.description",
    "amount": "i.amount",
    "comment": "i.comment",
}


def _apply_sort(sorts_allowed: dict, default: str = "date"):
    """Возвращает SQL-фрагмент ORDER BY из query-параметра sort=field:dir,field2:dir2.

    Пример: ?sort=amount:desc,date:asc
    """
    raw = request.args.get("sort", "").strip()
    if not raw:
        return f" ORDER BY {sorts_allowed[default]} DESC"
    parts = []
    for chunk in raw.split(","):
        chunk = chunk.strip()
        if not chunk:
            continue
        if ":" in chunk:
            field, direction = chunk.split(":", 1)
        else:
            field, direction = chunk, "asc"
        field = field.strip()
        direction = direction.strip().lower()
        if field not in sorts_allowed or direction not in ("asc", "desc"):
            continue
        sql_field = sorts_allowed[field]
        parts.append(f"{sql_field} {direction.upper()} NULLS LAST")
    if not parts:
        return f" ORDER BY {sorts_allowed[default]} DESC"
    return " ORDER BY " + ", ".join(parts)


def _apply_pagination():
    """Возвращает (limit, offset, page, page_size) из query ?page=1&page_size=50."""
    page = _parse_int(request.args.get("page"), 1) or 1
    if page < 1:
        page = 1
    page_size = _parse_int(request.args.get("page_size"), 50) or 50
    page_size = max(1, min(page_size, 500))
    return page_size, (page - 1) * page_size, page, page_size


# ============================================================
# EXPENSES
# ============================================================
@app.get("/api/expenses")
def list_expenses():
    months = _parse_months_param()
    where = ["TRUE"]
    p: list = []

    if months:
        where.append("EXTRACT(MONTH FROM e.date) = ANY(%s)")
        p.append(months)

    # ── категория / подкатегория ──
    cat_id = _parse_int(request.args.get("category_id"))
    cat_name = request.args.get("category", "").strip()
    if cat_id:
        where.append("e.category_id = %s"); p.append(cat_id)
    elif cat_name:
        where.append("ec.name = %s"); p.append(cat_name)

    sub_id = _parse_int(request.args.get("subcategory_id"))
    sub_name = request.args.get("subcategory", "").strip()
    if sub_id:
        where.append("e.subcategory_id = %s"); p.append(sub_id)
    elif sub_name:
        where.append("es.name = %s"); p.append(sub_name)

    # ── описание / комментарий (contains) ──
    desc = request.args.get("description", "").strip()
    if desc:
        where.append("e.description ILIKE %s ESCAPE '\\'")
        p.append("%" + _escape_like(desc) + "%")

    comment = request.args.get("comment", "").strip()
    if comment:
        where.append("COALESCE(e.comment,'') ILIKE %s ESCAPE '\\'")
        p.append("%" + _escape_like(comment) + "%")

    # ── дата (умно) ──
    _apply_date_filter(
        where, p, "e.date",
        date_from=request.args.get("date_from", "").strip(),
        date_to=request.args.get("date_to", "").strip(),
        date_eq=request.args.get("date", "").strip(),
    )

    # ── сумма (умно) ──
    _apply_amount_filter(
        where, p, "e.amount",
        amount_min=_parse_float(request.args.get("amount_min")),
        amount_max=_parse_float(request.args.get("amount_max")),
        amount_eq=_parse_float(request.args.get("amount")),
    )

    # ── quick search ──
    search = request.args.get("search", "").strip()
    if search:
        like = "%" + _escape_like(search) + "%"
        where.append(
            "(e.description ILIKE %s ESCAPE '\\' "
            "OR COALESCE(e.comment,'') ILIKE %s ESCAPE '\\' "
            "OR ec.name ILIKE %s ESCAPE '\\' "
            "OR es.name ILIKE %s ESCAPE '\\')"
        )
        p.extend([like, like, like, like])

    where_sql = " AND ".join(where)
    order_sql = _apply_sort(EXPENSE_SORTS)
    page_size, offset, page, _ = _apply_pagination()

    count_sql = f"""
        SELECT COUNT(*) AS c, COALESCE(SUM(e.amount),0) AS s
        FROM expenses e
        LEFT JOIN expense_categories ec    ON ec.id = e.category_id
        LEFT JOIN expense_subcategories es ON es.id = e.subcategory_id
        WHERE {where_sql}
    """
    cnt_row = q(count_sql, p, fetch="one")
    total = int(cnt_row["c"])
    total_amount = float(cnt_row["s"])

    data_sql = f"""
        SELECT e.id, e.date, e.description, e.amount, e.comment,
               ec.id AS category_id, ec.name AS category,
               es.id AS subcategory_id, es.name AS subcategory
        FROM expenses e
        LEFT JOIN expense_categories ec    ON ec.id = e.category_id
        LEFT JOIN expense_subcategories es ON es.id = e.subcategory_id
        WHERE {where_sql}
        {order_sql}
        LIMIT %s OFFSET %s
    """
    rows = q(data_sql, p + [page_size, offset])

    return jsonify({
        "items": ser(rows),
        "total": total,
        "total_amount": total_amount,
        "page": page,
        "page_size": page_size,
    })

# ============================================================
# INCOMES
# ============================================================
@app.get("/api/incomes")
def list_incomes():
    months = _parse_months_param()
    where = ["TRUE"]
    p: list = []

    if months:
        where.append("EXTRACT(MONTH FROM i.date) = ANY(%s)")
        p.append(months)

    cat_id = _parse_int(request.args.get("category_id"))
    cat_name = request.args.get("category", "").strip()
    if cat_id:
        where.append("i.category_id = %s"); p.append(cat_id)
    elif cat_name:
        where.append("ic.name = %s"); p.append(cat_name)

    desc = request.args.get("description", "").strip()
    if desc:
        where.append("i.description ILIKE %s ESCAPE '\\'")
        p.append("%" + _escape_like(desc) + "%")

    comment = request.args.get("comment", "").strip()
    if comment:
        where.append("COALESCE(i.comment,'') ILIKE %s ESCAPE '\\'")
        p.append("%" + _escape_like(comment) + "%")

    _apply_date_filter(
        where, p, "i.date",
        date_from=request.args.get("date_from", "").strip(),
        date_to=request.args.get("date_to", "").strip(),
        date_eq=request.args.get("date", "").strip(),
    )

    _apply_amount_filter(
        where, p, "i.amount",
        amount_min=_parse_float(request.args.get("amount_min")),
        amount_max=_parse_float(request.args.get("amount_max")),
        amount_eq=_parse_float(request.args.get("amount")),
    )

    search = request.args.get("search", "").strip()
    if search:
        like = "%" + _escape_like(search) + "%"
        where.append(
            "(i.description ILIKE %s ESCAPE '\\' "
            "OR COALESCE(i.comment,'') ILIKE %s ESCAPE '\\' "
            "OR ic.name ILIKE %s ESCAPE '\\')"
        )
        p.extend([like, like, like])

    where_sql = " AND ".join(where)
    order_sql = _apply_sort(INCOME_SORTS)
    page_size, offset, page, _ = _apply_pagination()

    count_sql = f"""
        SELECT COUNT(*) AS c, COALESCE(SUM(i.amount),0) AS s
        FROM incomes i
        LEFT JOIN income_categories ic ON ic.id = i.category_id
        WHERE {where_sql}
    """
    cnt_row = q(count_sql, p, fetch="one")
    total = int(cnt_row["c"])
    total_amount = float(cnt_row["s"])

    data_sql = f"""
        SELECT i.id, i.date, i.description, i.amount, i.comment,
               ic.id AS category_id, ic.name AS category
        FROM incomes i
        LEFT JOIN income_categories ic ON ic.id = i.category_id
        WHERE {where_sql}
        {order_sql}
        LIMIT %s OFFSET %s
    """
    rows = q(data_sql, p + [page_size, offset])

    return jsonify({
        "items": ser(rows),
        "total": total,
        "total_amount": total_amount,
        "page": page,
        "page_size": page_size,
    })

@app.post("/api/expenses")
def create_expense():
    d = request.get_json(force=True)
    row = q(
        "INSERT INTO expenses(date, category_id, subcategory_id, description, amount, comment) "
        "VALUES(%s, %s, %s, %s, %s, %s) RETURNING id",
        (d["date"], d.get("category_id"), d.get("subcategory_id"),
         d.get("description"), d["amount"], d.get("comment")),
        fetch="one",
    )
    return jsonify({"id": row["id"]}), 201


@app.put("/api/expenses/<int:eid>")
def update_expense(eid):
    d = request.get_json(force=True)
    q(
        "UPDATE expenses SET date=%s, category_id=%s, subcategory_id=%s, "
        "description=%s, amount=%s, comment=%s WHERE id=%s",
        (d["date"], d.get("category_id"), d.get("subcategory_id"),
         d.get("description"), d["amount"], d.get("comment"), eid),
    )
    return jsonify({"ok": True})


@app.delete("/api/expenses/<int:eid>")
def delete_expense(eid):
    q("DELETE FROM expenses WHERE id=%s", (eid,))
    return jsonify({"ok": True})

@app.post("/api/incomes")
def create_income():
    d = request.get_json(force=True)
    row = q(
        "INSERT INTO incomes(date, category_id, description, amount, comment) "
        "VALUES(%s, %s, %s, %s, %s) RETURNING id",
        (d["date"], d.get("category_id"), d.get("description"),
         d["amount"], d.get("comment")),
        fetch="one",
    )
    return jsonify({"id": row["id"]}), 201


@app.put("/api/incomes/<int:iid>")
def update_income(iid):
    d = request.get_json(force=True)
    q(
        "UPDATE incomes SET date=%s, category_id=%s, description=%s, amount=%s, comment=%s "
        "WHERE id=%s",
        (d["date"], d.get("category_id"), d.get("description"),
         d["amount"], d.get("comment"), iid),
    )
    return jsonify({"ok": True})


@app.delete("/api/incomes/<int:iid>")
def delete_income(iid):
    q("DELETE FROM incomes WHERE id=%s", (iid,))
    return jsonify({"ok": True})


# ============================================================
# REFS
# ============================================================
@app.get("/api/refs/expense-categories")
def list_exp_cats():
    return jsonify(ser(q("SELECT id, name FROM expense_categories ORDER BY name")))


@app.post("/api/refs/expense-categories")
def add_exp_cat():
    name = (request.get_json(force=True).get("name") or "").strip()
    if not name:
        abort(400)
    try:
        row = q(
            "INSERT INTO expense_categories(name) VALUES(%s) RETURNING id, name",
            (name,), fetch="one",
        )
        return jsonify(ser(row)), 201
    except errors.UniqueViolation:
        abort(409, "Такое имя уже есть")


@app.put("/api/refs/expense-categories/<int:cid>")
def rename_exp_cat(cid):
    name = (request.get_json(force=True).get("name") or "").strip()
    if not name:
        abort(400)
    try:
        q("UPDATE expense_categories SET name=%s WHERE id=%s", (name, cid))
        return jsonify({"ok": True})
    except errors.UniqueViolation:
        abort(409, "Такое имя уже есть")


@app.delete("/api/refs/expense-categories/<int:cid>")
def del_exp_cat(cid):
    q("DELETE FROM expense_categories WHERE id=%s", (cid,))
    return jsonify({"ok": True})


@app.get("/api/refs/expense-subcategories")
def list_subs():
    cid = request.args.get("category_id")
    if cid:
        return jsonify(ser(q(
            "SELECT id, category_id, name FROM expense_subcategories "
            "WHERE category_id=%s ORDER BY name", (int(cid),)
        )))
    return jsonify(ser(q(
        "SELECT id, category_id, name FROM expense_subcategories ORDER BY name"
    )))


@app.post("/api/refs/expense-subcategories")
def add_sub():
    d = request.get_json(force=True)
    cid, name = d.get("category_id"), (d.get("name") or "").strip()
    if not cid or not name:
        abort(400)
    try:
        row = q(
            "INSERT INTO expense_subcategories(category_id, name) "
            "VALUES(%s, %s) RETURNING id, category_id, name",
            (cid, name), fetch="one",
        )
        return jsonify(ser(row)), 201
    except errors.UniqueViolation:
        abort(409, "Такое имя уже есть")


@app.put("/api/refs/expense-subcategories/<int:sid>")
def rename_sub(sid):
    name = (request.get_json(force=True).get("name") or "").strip()
    if not name:
        abort(400)
    try:
        q("UPDATE expense_subcategories SET name=%s WHERE id=%s", (name, sid))
        return jsonify({"ok": True})
    except errors.UniqueViolation:
        abort(409, "Такое имя уже есть")


@app.delete("/api/refs/expense-subcategories/<int:sid>")
def del_sub(sid):
    q("DELETE FROM expense_subcategories WHERE id=%s", (sid,))
    return jsonify({"ok": True})


@app.get("/api/refs/income-categories")
def list_inc_cats():
    return jsonify(ser(q("SELECT id, name FROM income_categories ORDER BY name")))


@app.post("/api/refs/income-categories")
def add_inc_cat():
    name = (request.get_json(force=True).get("name") or "").strip()
    if not name:
        abort(400)
    try:
        row = q(
            "INSERT INTO income_categories(name) VALUES(%s) RETURNING id, name",
            (name,), fetch="one",
        )
        return jsonify(ser(row)), 201
    except errors.UniqueViolation:
        abort(409, "Такое имя уже есть")


@app.put("/api/refs/income-categories/<int:cid>")
def rename_inc_cat(cid):
    name = (request.get_json(force=True).get("name") or "").strip()
    if not name:
        abort(400)
    try:
        q("UPDATE income_categories SET name=%s WHERE id=%s", (name, cid))
        return jsonify({"ok": True})
    except errors.UniqueViolation:
        abort(409, "Такое имя уже есть")


@app.delete("/api/refs/income-categories/<int:cid>")
def del_inc_cat(cid):
    q("DELETE FROM income_categories WHERE id=%s", (cid,))
    return jsonify({"ok": True})


# ============================================================
# STATS
# ============================================================
@app.get("/api/stats")
def stats():
    months = _parse_months_param()
    m_filter = "WHERE EXTRACT(MONTH FROM date) = ANY(%s)" if months else ""
    m_filter_exp = "WHERE EXTRACT(MONTH FROM e.date) = ANY(%s)" if months else ""
    m_filter_inc = "WHERE EXTRACT(MONTH FROM i.date) = ANY(%s)" if months else ""
    p = [months] if months else []

    months_exp = q(f"""
        SELECT EXTRACT(MONTH FROM date)::int AS m, SUM(amount) AS s
        FROM expenses {m_filter} GROUP BY m ORDER BY m
    """, p if months else None)

    months_inc = q(f"""
        SELECT EXTRACT(MONTH FROM date)::int AS m, SUM(amount) AS s
        FROM incomes {m_filter} GROUP BY m ORDER BY m
    """, p if months else None)

    cats = q(f"""
        SELECT ec.name, SUM(e.amount) AS s
        FROM expenses e JOIN expense_categories ec ON ec.id = e.category_id
        {m_filter_exp}
        GROUP BY ec.name ORDER BY s DESC LIMIT 15
    """, p if months else None)

    subs = q(f"""
        SELECT es.name, SUM(e.amount) AS s
        FROM expenses e JOIN expense_subcategories es ON es.id = e.subcategory_id
        {m_filter_exp}
        GROUP BY es.name ORDER BY s DESC LIMIT 15
    """, p if months else None)

    inc_cats = q(f"""
        SELECT ic.name, SUM(i.amount) AS s
        FROM incomes i JOIN income_categories ic ON ic.id = i.category_id
        {m_filter_inc}
        GROUP BY ic.name ORDER BY s DESC
    """, p if months else None)

    return jsonify({
        "expenses_by_month": ser(months_exp),
        "incomes_by_month": ser(months_inc),
        "expenses_by_category": ser(cats),
        "expenses_by_subcategory": ser(subs),
        "incomes_by_category": ser(inc_cats),
    })

# ============================================================
# EXPORT
# ============================================================
def _csv_response(rows, fieldnames, filename):
    buf = io.StringIO()
    buf.write("\ufeff")  # BOM для Excel
    w = csv.DictWriter(
        buf, fieldnames=fieldnames,
        delimiter=";", quoting=csv.QUOTE_MINIMAL,
    )
    w.writeheader()
    for r in rows:
        w.writerow({k: (r.get(k) if r.get(k) is not None else "") for k in fieldnames})
    data = buf.getvalue().encode("utf-8")
    return send_file(
        io.BytesIO(data),
        mimetype="text/csv; charset=utf-8",
        as_attachment=True,
        download_name=filename,
    )


def _xlsx_response(headers, rows, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for col_idx, _ in enumerate(headers, 1):
        ws.column_dimensions[ws.cell(1, col_idx).column_letter].width = 18
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(
        out,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


def _export_data(kind):
    """kind: expenses|incomes|refs"""
    if kind == "expenses":
        rows = q("""
            SELECT e.id, e.date, ec.name AS category, es.name AS subcategory,
                   e.description, e.amount, e.comment
            FROM expenses e
            LEFT JOIN expense_categories ec    ON ec.id = e.category_id
            LEFT JOIN expense_subcategories es ON es.id = e.subcategory_id
            ORDER BY e.date, e.id
        """)
        rows = ser(rows)
        fields = [
            ("date", "Дата"),
            ("category", "Категория"),
            ("subcategory", "Подкатегория"),
            ("description", "Описание"),
            ("amount", "Стоимость"),
            ("comment", "Комментарий"),
        ]
        filename = "expenses"
    elif kind == "incomes":
        rows = q("""
            SELECT i.id, i.date, ic.name AS category, i.description, i.amount, i.comment
            FROM incomes i LEFT JOIN income_categories ic ON ic.id = i.category_id
            ORDER BY i.date, i.id
        """)
        rows = ser(rows)
        fields = [
            ("date", "Дата"),
            ("category", "Категория"),
            ("description", "Описание"),
            ("amount", "Сумма"),
            ("comment", "Комментарий"),
        ]
        filename = "incomes"
    else:  # refs
        cats = q("""
            SELECT ec.name AS category, es.name AS subcategory
            FROM expense_categories ec
            LEFT JOIN expense_subcategories es ON es.category_id = ec.id
            ORDER BY ec.name, es.name
        """)
        rows = [{"category": r["category"], "subcategory": r["subcategory"] or ""}
                for r in cats]
        incs = q("SELECT name FROM income_categories ORDER BY name")
        for r in incs:
            rows.append({"category": r["name"], "subcategory": "(доход)"})
        fields = [
            ("category", "Категория"),
            ("subcategory", "Подкатегория"),
        ]
        filename = "refs"
    return rows, fields, filename


@app.get("/api/export/<kind>.<fmt>")
def export_data(kind, fmt):
    if kind not in ("expenses", "incomes", "refs"):
        abort(404)
    if fmt not in ("csv", "xlsx"):
        abort(400)
    rows, fields, base = _export_data(kind)
    if fmt == "csv":
        return _csv_response(rows, [k for k, _ in fields], f"{base}.csv")
    headers = [v for _, v in fields]
    matrix = [[r.get(k) for k, _ in fields] for r in rows]
    return _xlsx_response(headers, matrix, f"{base}.xlsx")


# ============================================================
# IMPORT CSV
# ============================================================
def _detect_delimiter(sample: str) -> str:
    first = sample.splitlines()[0] if sample else ""
    return ";" if first.count(";") >= first.count(",") else ","


def _parse_date(s):
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(str(s).strip(), fmt).date()
        except Exception:
            pass
    raise ValueError(f"Неверная дата: {s}")


def _parse_amount(s):
    if s is None:
        return 0.0
    return float(str(s).replace(" ", "").replace(",", ".").strip() or 0)


def _human_error(exc: Exception) -> str:
    """Человеческое описание ошибки импорта."""
    if isinstance(exc, ValueError):
        return str(exc)
    if isinstance(exc, errors.UniqueViolation):
        return "Такая запись уже есть в базе (дубликат)"
    if isinstance(exc, errors.NotNullViolation):
        return "Не заполнено обязательное поле"
    if isinstance(exc, errors.ForeignKeyViolation):
        return "Ссылка на несуществующую категорию/подкатегорию"
    if isinstance(exc, errors.StringDataRightTruncation):
        return "Слишком длинное значение"
    if isinstance(exc, psycopg2.DataError):
        return "Неверный формат данных (проверьте дату и сумму)"
    if isinstance(exc, psycopg2.IntegrityError):
        return "Нарушение целостности данных"
    return f"Ошибка: {exc.__class__.__name__}"

@app.post("/api/import/expenses")
def import_expenses():
    if "file" not in request.files:
        abort(400, "Файл не передан")
    f = request.files["file"]
    raw = f.read()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = raw.decode("cp1251")
        except UnicodeDecodeError:
            return jsonify({
                "imported": 0, "skipped": 0, "total_errors": 1,
                "errors": ["Файл не в кодировке UTF-8 или Windows-1251"],
            }), 400

    if not text.strip():
        return jsonify({
            "imported": 0, "skipped": 0, "total_errors": 1,
            "errors": ["Файл пустой"],
        }), 400

    delim = _detect_delimiter(text)
    reader = csv.DictReader(io.StringIO(text), delimiter=delim)

    # Проверка обязательных колонок
    required = {"date", "category", "amount"}
    if reader.fieldnames:
        norm = {(_aliases_exp().get(k.strip().lower()) or k.strip().lower()) for k in reader.fieldnames}
        missing = required - norm
        if missing:
            return jsonify({
                "imported": 0, "skipped": 0, "total_errors": 1,
                "errors": [
                    f"В CSV не хватает колонок: {', '.join(sorted(missing))}. "
                    f"Ожидаются: дата, категория, стоимость, "
                    f"опционально — подкатегория, описание, комментарий."
                ],
            }), 400

    errors_list, imported, skipped = [], 0, 0
    conn = get_conn()
    try:
        with conn.cursor() as cur:
            for i, row in enumerate(reader, 2):
                try:
                    row2 = {_aliases_exp()[k.strip().lower()]: v
                            for k, v in row.items() if k and k.strip()}

                    if not row2.get("date"):
                        raise ValueError("не указана дата")
                    if not row2.get("category"):
                        raise ValueError("не указана категория")

                    d = _parse_date(row2["date"])
                    cat = row2["category"].strip()
                    if not cat:
                        raise ValueError("пустая категория")

                    cur.execute(
                        "INSERT INTO expense_categories(name) VALUES(%s) "
                        "ON CONFLICT(name) DO UPDATE SET name=EXCLUDED.name RETURNING id",
                        (cat,),
                    )
                    cid = cur.fetchone()[0]

                    sid = None
                    sub = (row2.get("subcategory") or "").strip()
                    if sub:
                        cur.execute(
                            "INSERT INTO expense_subcategories(category_id, name) "
                            "VALUES(%s, %s) "
                            "ON CONFLICT(category_id, name) DO UPDATE SET name=EXCLUDED.name "
                            "RETURNING id",
                            (cid, sub),
                        )
                        sid = cur.fetchone()[0]

                    amount = _parse_amount(row2.get("amount"))
                    desc = (row2.get("description") or "").strip() or None
                    comm = (row2.get("comment") or "").strip() or None

                    # Проверка дубликата
                    cur.execute(
                        "SELECT 1 FROM expenses WHERE date=%s AND category_id=%s "
                        "AND COALESCE(subcategory_id,-1)=COALESCE(%s,-1) "
                        "AND amount=%s AND COALESCE(description,'')=COALESCE(%s,'') "
                        "AND COALESCE(comment,'')=COALESCE(%s,'') LIMIT 1",
                        (d, cid, sid, amount, desc, comm),
                    )
                    if cur.fetchone():
                        skipped += 1
                        continue

                    cur.execute(
                        "INSERT INTO expenses(date, category_id, subcategory_id, "
                        "description, amount, comment) "
                        "VALUES(%s, %s, %s, %s, %s, %s)",
                        (d, cid, sid, desc, amount, comm),
                    )
                    imported += 1
                except Exception as e:
                    errors_list.append(f"Строка {i}: {_human_error(e)}")
            conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()

    return jsonify({
        "imported": imported,
        "skipped": skipped,
        "errors": errors_list[:30],
        "total_errors": len(errors_list),
    })


@app.post("/api/import/incomes")
def import_incomes():
    if "file" not in request.files:
        abort(400, "Файл не передан")
    f = request.files["file"]
    raw = f.read()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = raw.decode("cp1251")
        except UnicodeDecodeError:
            return jsonify({
                "imported": 0, "skipped": 0, "total_errors": 1,
                "errors": ["Файл не в кодировке UTF-8 или Windows-1251"],
            }), 400

    if not text.strip():
        return jsonify({
            "imported": 0, "skipped": 0, "total_errors": 1,
            "errors": ["Файл пустой"],
        }), 400

    delim = _detect_delimiter(text)
    reader = csv.DictReader(io.StringIO(text), delimiter=delim)

    required = {"date", "category", "amount"}
    if reader.fieldnames:
        norm = {(_aliases_inc().get(k.strip().lower()) or k.strip().lower()) for k in reader.fieldnames}
        missing = required - norm
        if missing:
            return jsonify({
                "imported": 0, "skipped": 0, "total_errors": 1,
                "errors": [
                    f"В CSV не хватает колонок: {', '.join(sorted(missing))}. "
                    f"Ожидаются: дата, категория, сумма, "
                    f"опционально — описание, комментарий."
                ],
            }), 400

    errors_list, imported, skipped = [], 0, 0
    conn = get_conn()
    try:
        with conn.cursor() as cur:
            for i, row in enumerate(reader, 2):
                try:
                    row2 = {_aliases_inc()[k.strip().lower()]: v
                            for k, v in row.items() if k and k.strip()}

                    if not row2.get("date"):
                        raise ValueError("не указана дата")
                    if not row2.get("category"):
                        raise ValueError("не указана категория")

                    d = _parse_date(row2["date"])
                    cat = row2["category"].strip()
                    if not cat:
                        raise ValueError("пустая категория")

                    cur.execute(
                        "INSERT INTO income_categories(name) VALUES(%s) "
                        "ON CONFLICT(name) DO UPDATE SET name=EXCLUDED.name RETURNING id",
                        (cat,),
                    )
                    cid = cur.fetchone()[0]

                    amount = _parse_amount(row2.get("amount"))
                    desc = (row2.get("description") or "").strip() or None
                    comm = (row2.get("comment") or "").strip() or None

                    cur.execute(
                        "SELECT 1 FROM incomes WHERE date=%s AND category_id=%s "
                        "AND amount=%s AND COALESCE(description,'')=COALESCE(%s,'') "
                        "AND COALESCE(comment,'')=COALESCE(%s,'') LIMIT 1",
                        (d, cid, amount, desc, comm),
                    )
                    if cur.fetchone():
                        skipped += 1
                        continue

                    cur.execute(
                        "INSERT INTO incomes(date, category_id, description, amount, comment) "
                        "VALUES(%s, %s, %s, %s, %s)",
                        (d, cid, desc, amount, comm),
                    )
                    imported += 1
                except Exception as e:
                    errors_list.append(f"Строка {i}: {_human_error(e)}")
            conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()

    return jsonify({
        "imported": imported,
        "skipped": skipped,
        "errors": errors_list[:30],
        "total_errors": len(errors_list),
    })

# ============================================================
# BACKUP DB
# ============================================================
@app.get("/api/backup")
def backup_db():
    url = os.environ["DATABASE_URL"]
    u = urlparse(url)
    env = os.environ.copy()
    env["PGPASSWORD"] = u.password or ""
    cmd = [
        "pg_dump",
        "-h", u.hostname or "db",
        "-p", str(u.port or 5432),
        "-U", u.username or "traty",
        "-d", (u.path or "/traty").lstrip("/"),
        "--no-owner", "--no-acl",
    ]
    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".sql", delete=False) as tmp:
            tmp_path = tmp.name
        with open(tmp_path, "wb") as fh:
            r = subprocess.run(cmd, stdout=fh, stderr=subprocess.PIPE, env=env)
        if r.returncode != 0:
            raise RuntimeError(r.stderr.decode("utf-8", "ignore"))
        return send_file(
            tmp_path,
            mimetype="application/sql",
            as_attachment=True,
            download_name=f"traty_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.sql",
        )
    except FileNotFoundError:
        abort(500, "pg_dump не установлен в контейнере")
    except Exception as e:
        abort(500, str(e))
    finally:
        # send_file может читать файл ещё в момент ответа — удаление
        # откладывается, т.к. Flask отдаёт файл потоком. Просто оставляем
        # его в /tmp, ОС почистит. (Корректнее: after_request, но для
        # локального инструмента допустимо.)
        pass

@app.get("/")
def root():
    index_path = os.path.join(FRONTEND_DIST, "index.html")
    if not os.path.exists(index_path):
        return (
            "React-фронт не собран. Соберите: cd frontend && npm install && npm run build",
            500,
        )
    return send_from_directory(FRONTEND_DIST, "index.html")


@app.get("/<path:path>")
def spa(path):
    """SPA-fallback: сначала статика из dist/, иначе index.html."""
    candidate = os.path.join(FRONTEND_DIST, path)
    if path and os.path.isfile(candidate):
        return send_from_directory(FRONTEND_DIST, path)
    index_path = os.path.join(FRONTEND_DIST, "index.html")
    if os.path.exists(index_path):
        return send_from_directory(FRONTEND_DIST, "index.html")
    abort(404)


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=3000, debug=True)