"""Импорт данных из Excel в PostgreSQL. Идемпотентно: пропускает, если БД не пуста."""
import os
import time
from datetime import datetime, date

import psycopg2
from openpyxl import load_workbook


def _wait_db(url: str, attempts: int = 30) -> None:
    for i in range(attempts):
        try:
            psycopg2.connect(url).close()
            return
        except psycopg2.OperationalError:
            print(f"[seed] waiting for DB ({i + 1}/{attempts})…")
            time.sleep(1)
    raise RuntimeError("DB not available")


def _as_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
            try:
                return datetime.strptime(v, fmt).date()
            except ValueError:
                pass
    return None


def _as_float(v) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def run() -> None:
    url = os.environ["DATABASE_URL"]
    path = os.environ.get("EXCEL_PATH", "data/траты.xlsx")
    _wait_db(url)

    conn = psycopg2.connect(url)
    conn.autocommit = False
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT COUNT(*) FROM expenses")
            if cur.fetchone()[0] > 0:
                print("[seed] DB already populated — skip")
                return
            cur.execute("SELECT COUNT(*) FROM incomes")
            if cur.fetchone()[0] > 0:
                print("[seed] DB already populated — skip")
                return

        if not os.path.exists(path):
            print(f"[seed] Excel not found at {path} — start with empty DB")
            return

        print(f"[seed] reading {path}…")
        wb = load_workbook(path, data_only=True, read_only=False)

        # ---------- Справочники ----------
        ws = wb["Справочники"]
        cat_map: dict[str, int] = {}
        sub_map: dict[tuple[int, str], int] = {}

        with conn.cursor() as cur:
            # категории расходов
            for c in range(1, 21):
                v = ws.cell(1, c).value
                if not v:
                    continue
                name = str(v).strip()
                if not name:
                    continue
                cur.execute(
                    "INSERT INTO expense_categories(name) VALUES(%s) "
                    "ON CONFLICT(name) DO UPDATE SET name=EXCLUDED.name RETURNING id",
                    (name,),
                )
                cid = cur.fetchone()[0]
                cat_map[name] = cid
                for r in range(2, ws.max_row + 1):
                    sv = ws.cell(r, c).value
                    if sv:
                        sub = str(sv).strip()
                        if sub:
                            cur.execute(
                                "INSERT INTO expense_subcategories(category_id, name) VALUES(%s, %s) "
                                "ON CONFLICT(category_id, name) DO NOTHING RETURNING id",
                                (cid, sub),
                            )
                            got = cur.fetchone()
                            if got:
                                sub_map[(cid, sub)] = got[0]
                            else:
                                cur.execute(
                                    "SELECT id FROM expense_subcategories WHERE category_id=%s AND name=%s",
                                    (cid, sub),
                                )
                                sub_map[(cid, sub)] = cur.fetchone()[0]

            # категории доходов
            for r in range(2, ws.max_row + 1):
                v = ws.cell(r, 21).value
                if v:
                    name = str(v).strip()
                    if name:
                        cur.execute(
                            "INSERT INTO income_categories(name) VALUES(%s) "
                            "ON CONFLICT(name) DO NOTHING",
                            (name,),
                        )
            conn.commit()

            cur.execute("SELECT id, name FROM income_categories")
            inc_cat_map = {name: cid for cid, name in cur.fetchall()}

        # ---------- Расходы ----------
        ws = wb["Учет расходов"]
        with conn.cursor() as cur:
            n = 0
            for r in range(2, ws.max_row + 1):
                d = _as_date(ws.cell(r, 3).value)
                raw_cat = ws.cell(r, 4).value
                if not d or not raw_cat:
                    continue
                cat = str(raw_cat).strip()
                if not cat:
                    continue

                cid = cat_map.get(cat)
                if not cid:
                    cur.execute(
                        "INSERT INTO expense_categories(name) VALUES(%s) RETURNING id",
                        (cat,),
                    )
                    cid = cur.fetchone()[0]
                    cat_map[cat] = cid

                sub = ws.cell(r, 5).value
                sub = str(sub).strip() if sub else None
                sid = None
                if sub:
                    key = (cid, sub)
                    if key in sub_map:
                        sid = sub_map[key]
                    else:
                        cur.execute(
                            "INSERT INTO expense_subcategories(category_id, name) "
                            "VALUES(%s, %s) RETURNING id",
                            (cid, sub),
                        )
                        sid = cur.fetchone()[0]
                        sub_map[key] = sid

                cur.execute(
                    "INSERT INTO expenses(date, category_id, subcategory_id, description, amount, comment) "
                    "VALUES(%s, %s, %s, %s, %s, %s)",
                    (
                        d, cid, sid,
                        (str(ws.cell(r, 6).value).strip() if ws.cell(r, 6).value else None),
                        _as_float(ws.cell(r, 7).value),
                        (str(ws.cell(r, 8).value).strip() if ws.cell(r, 8).value else None),
                    ),
                )
                n += 1
            conn.commit()
            print(f"[seed] imported {n} expenses")

        # ---------- Доходы ----------
        ws = wb["Учет доходов"]
        with conn.cursor() as cur:
            n = 0
            for r in range(2, ws.max_row + 1):
                d = _as_date(ws.cell(r, 3).value)
                raw_cat = ws.cell(r, 4).value
                if not d or not raw_cat:
                    continue
                cat = str(raw_cat).strip()
                if not cat:
                    continue
                cid = inc_cat_map.get(cat)
                if not cid:
                    cur.execute(
                        "INSERT INTO income_categories(name) VALUES(%s) RETURNING id",
                        (cat,),
                    )
                    cid = cur.fetchone()[0]
                    inc_cat_map[cat] = cid

                cur.execute(
                    "INSERT INTO incomes(date, category_id, description, amount, comment) "
                    "VALUES(%s, %s, %s, %s, %s)",
                    (
                        d, cid,
                        (str(ws.cell(r, 5).value).strip() if ws.cell(r, 5).value else None),
                        _as_float(ws.cell(r, 6).value),
                        (str(ws.cell(r, 9).value).strip() if ws.cell(r, 9).value else None),
                    ),
                )
                n += 1
            conn.commit()
            print(f"[seed] imported {n} incomes")

        print("[seed] done ✔")
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()